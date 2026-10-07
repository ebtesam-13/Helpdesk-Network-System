import os
import sqlite3
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------
# 1. مسار قاعدة البيانات المشتركة على الشبكة
# ---------------------------------------------------------
DB_PATH = r"\\DESKTOP-CK5V925\HelpdeskShare\tickets.db"

# ---------------------------------------------------------
# 2. نصوص اللغتين (العربية والإنجليزية)
# ---------------------------------------------------------
LANGUAGES = {
    "ar": {
        "title": "🛠️ لوحة تحكم الدعم الفني (IT Admin)",
        "switch_lang": "English 🌐",
        "search_btn": "🔍 بحث",
        "reset_btn": "🔄 تحديث البيانات",
        "export_btn": "📊 تصدير Excel",
        "col_id": "رقم التذكرة",
        "col_emp": "اسم الموظف",
        "col_issue": "المشكلة",
        "col_priority": "الأولوية",
        "col_status": "الحالة",
        "col_action": "الإجراء (انقر مرتين لإغلاق)",
        "msg_error_db": "تعذر الاتصال بقاعدة البيانات المشتركة:\n",
        "msg_confirm_close": "هل أنت تأكد من تغيير حالة التذكرة رقم {id} إلى (تم الحل)؟",
        "msg_resolve_success": "تم إغلاق التذكرة بنجاح!",
        "msg_export_success": "تم تصدير التقرير بنجاح إلى ملف Excel:\n",
        "status_open": "مفتوحة",
        "status_closed": "تم الحل",
    },
    "en": {
        "title": "🛠️ IT Support Dashboard (Admin)",
        "switch_lang": "العربية 🌐",
        "search_btn": "🔍 Search",
        "reset_btn": "🔄 Refresh Data",
        "export_btn": "📊 Export Excel",
        "col_id": "Ticket ID",
        "col_emp": "Employee",
        "col_issue": "Issue",
        "col_priority": "Priority",
        "col_status": "Status",
        "col_action": "Action (Double-click to close)",
        "msg_error_db": "Could not connect to shared database:\n",
        "msg_confirm_close": "Are you sure you want to resolve ticket #{id}?",
        "msg_resolve_success": "Ticket closed successfully!",
        "msg_export_success": "Report exported successfully to Excel:\n",
        "status_open": "Open",
        "status_closed": "Resolved",
    }
}

current_lang = "ar"

# ---------------------------------------------------------
# 3. الدوال البرمجية
# ---------------------------------------------------------
def init_db():
    try:
        db_dir = os.path.dirname(DB_PATH)
        if db_dir and not os.path.exists(db_dir) and not DB_PATH.startswith("\\\\"):
            os.makedirs(db_dir, exist_ok=True)

        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS tickets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee TEXT NOT NULL,
                issue TEXT NOT NULL,
                priority TEXT NOT NULL,
                status TEXT DEFAULT 'مفتوحة'
            )
        """)
        conn.commit()
        conn.close()
    except Exception as e:
        messagebox.showerror("خطأ / Error", f"{LANGUAGES[current_lang]['msg_error_db']}{e}")

def load_tickets(query=""):
    entry_search.delete(0, tk.END)
    for row in tree.get_children():
        tree.delete(row)

    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        if query:
            cursor.execute("""
                SELECT * FROM tickets 
                WHERE CAST(id AS TEXT) LIKE ? 
                   OR employee LIKE ? 
                   OR issue LIKE ? 
                   OR priority LIKE ? 
                   OR status LIKE ?
                ORDER BY id DESC
            """, (f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%", f"%{query}%"))
        else:
            cursor.execute("SELECT * FROM tickets ORDER BY id DESC")
            
        rows = cursor.fetchall()
        conn.close()

        for idx, row in enumerate(rows):
            tag = "even" if idx % 2 == 0 else "odd"
            action_text = "✅ إغلاق التذكرة" if current_lang == "ar" else "✅ Close Ticket"
            display_row = list(row) + [action_text]
            tree.insert("", tk.END, values=display_row, tags=(tag,))
    except Exception as e:
        messagebox.showerror("خطأ / Error", f"{LANGUAGES[current_lang]['msg_error_db']}{e}")

def on_row_double_click(event):
    selected_item = tree.selection()
    if not selected_item:
        return

    item_values = tree.item(selected_item)["values"]
    ticket_id = item_values[0]
    current_status = item_values[4]

    if current_status == "تم الحل":
        return

    confirm = messagebox.askyesno(
        "تأكيد / Confirm", 
        LANGUAGES[current_lang]["msg_confirm_close"].format(id=ticket_id)
    )

    if confirm:
        try:
            conn = sqlite3.connect(DB_PATH)
            cursor = conn.cursor()
            cursor.execute("UPDATE tickets SET status = 'تم الحل' WHERE id = ?", (ticket_id,))
            conn.commit()
            conn.close()

            messagebox.showinfo("تحديث / Update", LANGUAGES[current_lang]["msg_resolve_success"])
            load_tickets()
        except Exception as e:
            messagebox.showerror("خطأ / Error", f"{e}")

def search_tickets():
    query = entry_search.get().strip()
    load_tickets(query)

def export_to_excel():
    file_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")],
        title="حفظ التقرير كـ Excel"
    )

    if not file_path:
        return

    try:
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT id, employee, issue, priority, status FROM tickets ORDER BY id DESC")
        rows = cursor.fetchall()
        conn.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tickets Report"
        ws.sheet_view.rightToLeft = (current_lang == "ar")

        headers = [
            LANGUAGES[current_lang]["col_id"],
            LANGUAGES[current_lang]["col_emp"],
            LANGUAGES[current_lang]["col_issue"],
            LANGUAGES[current_lang]["col_priority"],
            LANGUAGES[current_lang]["col_status"]
        ]
        ws.append(headers)

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        align_center = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        for row in rows:
            ws.append(row)

        for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
            for cell in row:
                cell.alignment = align_center

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

        wb.save(file_path)
        messagebox.showinfo("تصدير / Export", f"{LANGUAGES[current_lang]['msg_export_success']}{file_path}")
    except Exception as e:
        messagebox.showerror("خطأ / Error", f"{e}")

def toggle_language():
    global current_lang
    current_lang = "en" if current_lang == "ar" else "ar"
    
    btn_lang.config(text=LANGUAGES[current_lang]["switch_lang"])
    lbl_title.config(text=LANGUAGES[current_lang]["title"])
    btn_search.config(text=LANGUAGES[current_lang]["search_btn"])
    btn_reset.config(text=LANGUAGES[current_lang]["reset_btn"])
    btn_export.config(text=LANGUAGES[current_lang]["export_btn"])

    tree.heading("id", text=LANGUAGES[current_lang]["col_id"])
    tree.heading("employee", text=LANGUAGES[current_lang]["col_emp"])
    tree.heading("issue", text=LANGUAGES[current_lang]["col_issue"])
    tree.heading("priority", text=LANGUAGES[current_lang]["col_priority"])
    tree.heading("status", text=LANGUAGES[current_lang]["col_status"])
    tree.heading("action", text=LANGUAGES[current_lang]["col_action"])

    load_tickets()

# ---------------------------------------------------------
# 4. بناء الواجهة الرسومية (GUI)
# ---------------------------------------------------------
init_db()

root = tk.Tk()
root.title("IT Support Helpdesk Admin")
root.geometry("950x650")
root.configure(bg="#F5F7FA")

FONT_FAMILY = "Segoe UI"

style = ttk.Style()
style.theme_use("clam")
style.configure("Treeview",
                background="#FFFFFF",
                foreground="#2D3748",
                rowheight=32,
                fieldbackground="#FFFFFF",
                font=(FONT_FAMILY, 10))
style.configure("Treeview.Heading",
                background="#1A365D",
                foreground="#FFFFFF",
                font=(FONT_FAMILY, 10, "bold"),
                padding=6)
style.map("Treeview.Heading", background=[('active', '#2B6CB0')])

header_frame = tk.Frame(root, bg="#1A365D", height=60)
header_frame.pack(fill="x", side="top")

btn_lang = tk.Button(header_frame, text=LANGUAGES[current_lang]["switch_lang"], command=toggle_language,
                     bg="#2B6CB0", fg="white", font=(FONT_FAMILY, 9, "bold"), relief="flat", padx=10, pady=3, cursor="hand2")
btn_lang.pack(side="right", padx=15, pady=15)

lbl_title = tk.Label(header_frame, text=LANGUAGES[current_lang]["title"], font=(FONT_FAMILY, 15, "bold"), bg="#1A365D", fg="#FFFFFF")
lbl_title.pack(pady=12)

# شريط البحث المتمركز في الأعلى
search_container = tk.Frame(root, bg="#F5F7FA")
search_container.pack(pady=15, fill="x")

search_box = tk.Frame(search_container, bg="#F5F7FA")
search_box.pack(anchor="center")

entry_search = tk.Entry(search_box, font=(FONT_FAMILY, 10), bg="#FFFFFF", relief="solid", bd=1, width=32, justify="center")
entry_search.pack(side="left", padx=5)

btn_search = tk.Button(search_box, text=LANGUAGES[current_lang]["search_btn"], command=search_tickets,
                       bg="#DD6B20", fg="white", font=(FONT_FAMILY, 9, "bold"), relief="flat", padx=12, pady=3, cursor="hand2")
btn_search.pack(side="left", padx=5)

btn_reset = tk.Button(search_box, text=LANGUAGES[current_lang]["reset_btn"], command=lambda: load_tickets(),
                      bg="#718096", fg="white", font=(FONT_FAMILY, 9), relief="flat", padx=10, pady=3, cursor="hand2")
btn_reset.pack(side="left", padx=5)

# منطقة جدول التذاكر
table_frame = tk.Frame(root, bg="#F5F7FA")
table_frame.pack(pady=5, padx=20, fill="both", expand=True)

columns = ("id", "employee", "issue", "priority", "status", "action")
tree = ttk.Treeview(table_frame, columns=columns, show="headings")

tree.heading("id", text=LANGUAGES[current_lang]["col_id"])
tree.heading("employee", text=LANGUAGES[current_lang]["col_emp"])
tree.heading("issue", text=LANGUAGES[current_lang]["col_issue"])
tree.heading("priority", text=LANGUAGES[current_lang]["col_priority"])
tree.heading("status", text=LANGUAGES[current_lang]["col_status"])
tree.heading("action", text=LANGUAGES[current_lang]["col_action"])

tree.column("id", width=80, anchor="center")
tree.column("employee", width=130, anchor="center")
tree.column("issue", width=240, anchor="center")
tree.column("priority", width=90, anchor="center")
tree.column("status", width=100, anchor="center")
tree.column("action", width=160, anchor="center")

tree.tag_configure("even", background="#FFFFFF")
tree.tag_configure("odd", background="#EDF2F7")

# ربط النقر المزدوج بإغلاق التذكرة
tree.bind("<Double-1>", on_row_double_click)

tree.pack(fill="both", expand=True)

# الشريط السفلي (زر تصدير Excel فقط في الأسفل يساراً)
bottom_frame = tk.Frame(root, bg="#F5F7FA")
bottom_frame.pack(fill="x", padx=20, pady=15)

btn_export = tk.Button(bottom_frame, text=LANGUAGES[current_lang]["export_btn"], command=export_to_excel,
                       bg="#38A169", fg="white", font=(FONT_FAMILY, 10, "bold"), relief="flat", padx=15, pady=6, cursor="hand2")
btn_export.pack(side="left")

load_tickets()
root.mainloop()